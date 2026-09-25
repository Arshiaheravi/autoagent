"""Tests for replicator.py — URL/image → design brief → handoff pipeline."""
import json
from pathlib import Path
from unittest.mock import MagicMock, patch

import pytest

import replicator


@pytest.fixture(autouse=True)
def _tmp_replicate_dir(tmp_path, monkeypatch):
    """Redirect REPLICATE_DIR into tmp_path per test."""
    monkeypatch.setattr(replicator, "REPLICATE_DIR", tmp_path / "replicate")
    (tmp_path / "replicate").mkdir(parents=True, exist_ok=True)
    yield


# ── Slugs ─────────────────────────────────────────────────────────────

def test_slug_from_url():
    slug = replicator._slug("https://example.com/pricing")
    assert slug.startswith("example-com-pricing-")


def test_slug_from_image_path():
    slug = replicator._slug("/tmp/my-design.png")
    assert slug.startswith("my-design-")


def test_slug_strips_unsafe_chars():
    slug = replicator._slug("https://example.com/foo/bar?q=baz")
    assert "?" not in slug
    assert "=" not in slug


def test_is_url():
    assert replicator._is_url("https://example.com")
    assert replicator._is_url("http://example.com")
    assert not replicator._is_url("/tmp/a.png")
    assert not replicator._is_url("a.png")


# ── Capture image ─────────────────────────────────────────────────────

def test_capture_image_copies_file(tmp_path):
    src = tmp_path / "ref.png"
    src.write_bytes(b"\x89PNG\r\n\x1a\n" + b"\x00" * 100)
    out_dir = tmp_path / "out"
    result = replicator.capture_image(str(src), out_dir)
    assert result["kind"] == "image"
    assert result["source"] == str(src)
    assert len(result["screenshots"]) == 1
    assert Path(result["screenshots"][0]).exists()
    assert (out_dir / "source.txt").read_text() == str(src)


def test_capture_image_missing_file(tmp_path):
    result = replicator.capture_image("/nonexistent/x.png", tmp_path / "out")
    assert "error" in result
    assert "not found" in result["error"].lower()


def test_capture_image_unsupported_ext(tmp_path):
    src = tmp_path / "ref.bmp"
    src.write_bytes(b"bmp")
    result = replicator.capture_image(str(src), tmp_path / "out")
    assert "error" in result
    assert "unsupported" in result["error"].lower()


# ── Capture dispatch ──────────────────────────────────────────────────

def test_capture_dispatches_image(tmp_path):
    src = tmp_path / "ref.png"
    src.write_bytes(b"\x89PNG\r\n\x1a\n" + b"\x00" * 100)
    with patch.object(replicator, "capture_url") as mock_url:
        result = replicator.capture(str(src), tmp_path / "out")
    mock_url.assert_not_called()
    assert result["kind"] == "image"


def test_capture_dispatches_url(tmp_path):
    with patch.object(replicator, "capture_url", return_value={"kind": "url"}) as mock_url:
        result = replicator.capture("https://example.com", tmp_path / "out")
    mock_url.assert_called_once()
    assert result["kind"] == "url"


# ── Spec extraction ───────────────────────────────────────────────────

SAMPLE_SPEC = {
    "movement": "Minimal Precision",
    "layout": {"structure": "stacked", "sections": ["hero", "features", "footer"]},
    "palette": {"bg": "#0a0a0a", "surface": "#141414", "accents": ["#00c896"], "text": "#ffffff"},
    "typography": {"display": "Inter", "body": "Inter", "mono": "JetBrains Mono", "scale_rule": "4xl/xl/base"},
    "components": [{"name": "Hero", "notes": "Centered headline + CTA"}],
    "motion": ["fade-up on scroll"],
    "copy": {"headline": "Ship fast", "subhead": "Stay sane", "cta": "Get started"},
    "assets_needed": ["logo"],
    "tech_stack_hint": "React+Tailwind",
    "notes": "Clean dark minimal.",
}


def test_extract_spec_propagates_capture_error():
    result = replicator.extract_spec({"error": "bad network"})
    assert result == {"error": "bad network"}


def test_extract_spec_calls_vision_and_attaches_metadata(tmp_path):
    png = tmp_path / "shot.png"
    png.write_bytes(b"\x89PNG\r\n\x1a\n" + b"\x00" * 10)
    capture_data = {
        "kind": "url",
        "source": "https://example.com",
        "screenshots": [str(png)],
        "dom_path": None,
        "meta": {},
    }
    with patch.object(replicator, "_call_claude_vision", return_value=dict(SAMPLE_SPEC)):
        spec = replicator.extract_spec(capture_data)
    assert spec["_source"] == "https://example.com"
    assert spec["_kind"] == "url"
    assert spec["movement"] == "Minimal Precision"


def test_extract_spec_handles_vision_failure(tmp_path):
    png = tmp_path / "shot.png"
    png.write_bytes(b"x")
    capture_data = {"kind": "image", "source": str(png), "screenshots": [str(png)], "dom_path": None, "meta": {}}
    with patch.object(replicator, "_call_claude_vision", side_effect=RuntimeError("api down")):
        spec = replicator.extract_spec(capture_data)
    assert "error" in spec
    assert "api down" in spec["error"]


# ── Design-system merge ───────────────────────────────────────────────

def test_merge_design_system_reads_css_tokens(tmp_path):
    frontend = tmp_path / "frontend"
    frontend.mkdir()
    (frontend / "styles.css").write_text(
        ":root {\n  --bg: #0a0a0a;\n  --accent: #00c896;\n  --text: #ffffff;\n}\n"
    )
    spec = dict(SAMPLE_SPEC)
    merged = replicator.merge_design_system(spec, tmp_path)
    assert merged["_project_tokens"]["bg"] == "#0a0a0a"
    assert merged["_project_tokens"]["accent"] == "#00c896"
    assert "styles.css" in merged["_design_system_source"]


def test_merge_design_system_no_css(tmp_path):
    spec = dict(SAMPLE_SPEC)
    merged = replicator.merge_design_system(spec, tmp_path)
    assert "_project_tokens" not in merged


def test_merge_design_system_nonexistent_project():
    spec = dict(SAMPLE_SPEC)
    merged = replicator.merge_design_system(spec, Path("/nonexistent"))
    assert merged == spec


# ── Brief writer ──────────────────────────────────────────────────────

def test_write_brief_generates_md_and_json(tmp_path):
    spec = dict(SAMPLE_SPEC)
    spec["_source"] = "https://example.com"
    spec["_kind"] = "url"
    out = replicator.write_brief(spec, tmp_path)
    assert out.name == "brief.md"
    assert out.exists()
    assert (tmp_path / "brief.json").exists()
    md = out.read_text(encoding="utf-8")
    assert "Minimal Precision" in md
    assert "#0a0a0a" in md
    assert "Ship fast" in md
    assert "Operator annotations" in md
    data = json.loads((tmp_path / "brief.json").read_text(encoding="utf-8"))
    assert data["movement"] == "Minimal Precision"


def test_write_brief_includes_project_tokens(tmp_path):
    spec = dict(SAMPLE_SPEC)
    spec["_project_tokens"] = {"bg": "#111", "accent": "#0f0"}
    out = replicator.write_brief(spec, tmp_path)
    md = out.read_text(encoding="utf-8")
    assert "Project design-system tokens" in md
    assert "--bg" in md


def test_write_brief_error_path(tmp_path):
    out = replicator.write_brief({"error": "spec failed"}, tmp_path)
    md = out.read_text(encoding="utf-8")
    assert "ERROR" in md
    assert "spec failed" in md


# ── Handoff ───────────────────────────────────────────────────────────

def test_handoff_pending_review_notifies_seb(tmp_path):
    brief = tmp_path / "brief.md"
    brief.write_text("# brief")
    with patch.object(replicator, "_notify_seb") as mock_notify:
        result = replicator.handoff(brief, wait_for_review=True)
    mock_notify.assert_called_once()
    assert result["status"] == "pending_review"
    assert result["brief_path"] == str(brief)


def test_handoff_no_wait_runs_council(tmp_path):
    brief = tmp_path / "brief.md"
    brief.write_text("# brief content")
    fake_council = MagicMock(return_value={"synthesis": "Looks good. Ship it."})
    with patch.dict("sys.modules", {"council": MagicMock(convene_council=fake_council)}):
        result = replicator.handoff(brief, wait_for_review=False)
    assert result["status"] == "ready"
    assert "Looks good" in result["council_synthesis"]


# ── Top-level replicate ───────────────────────────────────────────────

def test_replicate_failed_capture_returns_failed(tmp_path):
    with patch.object(replicator, "capture", return_value={"error": "net down"}):
        result = replicator.replicate("https://example.com")
    assert result["status"] == "failed"
    assert "net down" in result["error"]


def test_replicate_failed_spec_returns_failed(tmp_path):
    png = tmp_path / "r.png"
    png.write_bytes(b"\x89PNG\r\n\x1a\n" + b"\x00" * 10)
    with patch.object(replicator, "capture",
                      return_value={"kind": "image", "source": str(png), "screenshots": [str(png)],
                                    "dom_path": None, "meta": {}}):
        with patch.object(replicator, "extract_spec", return_value={"error": "vision fail"}):
            result = replicator.replicate(str(png))
    assert result["status"] == "failed"


def test_replicate_happy_path_pending_review(tmp_path):
    png = tmp_path / "r.png"
    png.write_bytes(b"\x89PNG\r\n\x1a\n" + b"\x00" * 10)
    fake_capture = {"kind": "image", "source": str(png), "screenshots": [str(png)],
                    "dom_path": None, "meta": {}}
    with patch.object(replicator, "capture", return_value=fake_capture), \
         patch.object(replicator, "extract_spec", return_value=dict(SAMPLE_SPEC, _source=str(png), _kind="image")), \
         patch.object(replicator, "a11y_review", return_value=""), \
         patch.object(replicator, "_notify_seb"):
        result = replicator.replicate(str(png))
    assert result["status"] == "pending_review"
    assert "slug" in result
    assert Path(result["brief_path"]).exists()


# ── PDF + DOCX capture ────────────────────────────────────────────────

def test_capture_pdf_copies_file(tmp_path):
    src = tmp_path / "ref.pdf"
    src.write_bytes(b"%PDF-1.4\n" + b"\x00" * 50)
    out_dir = tmp_path / "out"
    result = replicator.capture_pdf(str(src), out_dir)
    assert result["kind"] == "pdf"
    assert len(result["screenshots"]) == 1
    assert Path(result["screenshots"][0]).exists()


def test_capture_pdf_missing(tmp_path):
    result = replicator.capture_pdf("/nope.pdf", tmp_path / "out")
    assert "error" in result


def test_capture_docx_extracts_text(tmp_path):
    import docx as _docx
    src = tmp_path / "ref.docx"
    doc = _docx.Document()
    doc.add_paragraph("Hero headline: Ship fast")
    doc.add_paragraph("Dark theme, terminal aesthetic")
    doc.save(str(src))

    out_dir = tmp_path / "out"
    result = replicator.capture_docx(str(src), out_dir)
    assert result["kind"] == "docx"
    extracted = (out_dir / "extracted.txt").read_text(encoding="utf-8")
    assert "Ship fast" in extracted
    assert result["meta"]["paragraphs"] == 2
    assert result["screenshots"] == []


def test_capture_dispatch_pdf(tmp_path):
    pdf = tmp_path / "ref.pdf"
    pdf.write_bytes(b"%PDF-1.4\n")
    result = replicator.capture(str(pdf), tmp_path / "out")
    assert result["kind"] == "pdf"


def test_capture_dispatch_docx(tmp_path):
    import docx as _docx
    src = tmp_path / "ref.docx"
    doc = _docx.Document()
    doc.add_paragraph("test")
    doc.save(str(src))
    result = replicator.capture(str(src), tmp_path / "out")
    assert result["kind"] == "docx"


# ── Brand-asset scanner ───────────────────────────────────────────────

def test_scan_brand_assets_finds_logos_and_fonts(tmp_path):
    pub = tmp_path / "public"
    pub.mkdir()
    (pub / "logo.svg").write_text("<svg/>")
    (pub / "brand-mark.png").write_bytes(b"\x89PNG")
    (pub / "photo.jpg").write_bytes(b"jpg")
    fonts = tmp_path / "public" / "fonts"
    fonts.mkdir()
    (fonts / "Inter.woff2").write_bytes(b"woff")

    result = replicator._scan_brand_assets(tmp_path)
    assert any("logo.svg" in p for p in result["logos"])
    assert any("brand-mark.png" in p for p in result["logos"])
    assert any("photo.jpg" in p for p in result["images"])
    assert any("Inter.woff2" in p for p in result["fonts"])


def test_scan_brand_assets_empty_project(tmp_path):
    result = replicator._scan_brand_assets(tmp_path)
    assert result == {"logos": [], "fonts": [], "images": []}


def test_merge_design_system_adds_brand_assets(tmp_path):
    (tmp_path / "public").mkdir()
    (tmp_path / "public" / "logo.svg").write_text("<svg/>")
    spec = dict(SAMPLE_SPEC)
    merged = replicator.merge_design_system(spec, tmp_path)
    assert "_brand_assets" in merged
    assert len(merged["_brand_assets"]["logos"]) == 1


# ── Brief writer: brand + a11y sections ───────────────────────────────

def test_write_brief_includes_brand_section(tmp_path):
    spec = dict(SAMPLE_SPEC)
    spec["_brand_assets"] = {
        "logos": ["public/logo.svg"],
        "fonts": ["public/fonts/Inter.woff2"],
        "images": [],
    }
    out = replicator.write_brief(spec, tmp_path)
    md = out.read_text(encoding="utf-8")
    assert "Brand assets" in md
    assert "logo.svg" in md
    assert "Inter.woff2" in md


def test_write_brief_includes_a11y_section(tmp_path):
    spec = dict(SAMPLE_SPEC)
    spec["_a11y_review"] = "**Contrast:** fails 4.5:1 on muted text."
    out = replicator.write_brief(spec, tmp_path)
    md = out.read_text(encoding="utf-8")
    assert "Accessibility review" in md
    assert "fails 4.5:1" in md


# ── A11y review ───────────────────────────────────────────────────────

def test_a11y_review_returns_stdout(monkeypatch):
    class FakeResult:
        returncode = 0
        stdout = "**Contrast:** OK"
        stderr = ""
    monkeypatch.setattr(replicator.subprocess, "run", lambda *a, **k: FakeResult())
    out = replicator.a11y_review(SAMPLE_SPEC)
    assert "Contrast" in out


def test_a11y_review_tolerates_timeout(monkeypatch):
    def boom(*a, **k):
        raise replicator.subprocess.TimeoutExpired(cmd="claude", timeout=120)
    monkeypatch.setattr(replicator.subprocess, "run", boom)
    out = replicator.a11y_review(SAMPLE_SPEC)
    assert out == ""


# ── extract_spec text-only (docx path) ────────────────────────────────

def test_extract_spec_docx_uses_text_context(tmp_path):
    txt = tmp_path / "extracted.txt"
    txt.write_text("Headline: Ship fast\nDark terminal theme", encoding="utf-8")
    capture_data = {
        "kind": "docx",
        "source": "/some/ref.docx",
        "screenshots": [],
        "dom_path": str(txt),
        "meta": {"format": "docx"},
    }
    with patch.object(replicator, "_call_claude_vision", return_value=dict(SAMPLE_SPEC)) as mock_vis:
        spec = replicator.extract_spec(capture_data)
    assert spec["_kind"] == "docx"
    call_args = mock_vis.call_args
    assert call_args.args[0] == []  # no screenshots
    assert "Ship fast" in call_args.kwargs["extra_context"]
